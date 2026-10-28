"""Regenerate small real format fixtures. Not a CI/runtime dependency.

Requires Pillow, python-docx, openpyxl, python-pptx, LibreOffice and FFmpeg.
The production renderer/converter is never replaced by these fixture generators.
"""
import argparse
import io
import gzip
import json
import math
from pathlib import Path
import struct
import subprocess
import tarfile
import tempfile
import wave
import zipfile

from PIL import Image, ImageDraw
from docx import Document
from docx.shared import Inches
from openpyxl import Workbook
from pptx import Presentation
from pptx.util import Inches as SlideInches

ROOT = Path(__file__).resolve().parents[3]
DEST = ROOT / 'talos-e2e/fixtures/resource-preview'


def generate(ffmpeg):
    DEST.mkdir(parents=True, exist_ok=True)
    cases = []

    def add(name, mime, viewer, default=False, **expected):
        cases.append(dict(name=name, mime=mime, viewer=viewer, admission='default' if default else 'opt-in', **expected))

    def convert(source, target):
        with tempfile.TemporaryDirectory() as profile:
            subprocess.run(['soffice', '-env:UserInstallation=' + Path(profile).as_uri(), '--headless', '--convert-to', target,
                            '--outdir', str(DEST), str(DEST / source)], check=True, capture_output=True, timeout=90)
        result = DEST / (Path(source).stem + '.' + target.split(':')[0])
        assert result.is_file(), str(result)
        return result.name

    image = Image.new('RGB', (160, 96), '#126543')
    ImageDraw.Draw(image).text((12, 35), 'Talos Preview', fill='white')
    for ext, mime in [('png', 'image/png'), ('jpg', 'image/jpeg'), ('webp', 'image/webp'), ('gif', 'image/gif'), ('bmp', 'image/bmp')]:
        name = 'matrix-image.' + ext
        image.save(DEST / name)
        add(name, mime, 'image', ext in ('png', 'jpg', 'webp'), width=160, height=96)
    svg = 'matrix-image.svg'
    (DEST / svg).write_text('<svg xmlns="http://www.w3.org/2000/svg" width="160" height="96"><rect width="160" height="96" fill="#126543"/><text x="12" y="55" fill="white">Talos SVG</text></svg>')
    add(svg, 'text/plain', 'text', text=(DEST / svg).read_text(), note='Files detects SVG as plain text; Worker file metadata can supply image/svg+xml.')

    doc = Document()
    doc.add_heading('Word preview first page', 0)
    doc.add_paragraph('Chinese content: 中文预览。 Bold, tables and embedded images.')
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = 'Item'; table.cell(0, 1).text = 'Value'
    table.cell(1, 0).text = 'Preview total'; table.cell(1, 1).text = '42'
    doc.add_picture(str(DEST / 'matrix-image.png'), width=Inches(1.6))
    doc.add_page_break(); doc.add_heading('Word preview second page', 0)
    doc.save(DEST / 'matrix-word.docx')
    for name, mime, default in [('matrix-word.docx', 'application/vnd.openxmlformats-officedocument.wordprocessingml.document', True),
                              (convert('matrix-word.docx', 'doc:MS Word 97'), 'application/msword', True),
                              (convert('matrix-word.docx', 'odt'), 'application/vnd.oasis.opendocument.text', False),
                              (convert('matrix-word.docx', 'pdf'), 'application/pdf', True)]:
        add(name, mime, 'pdf', default, text='Word preview first page', secondText='Word preview second page', chineseText='中文预览')

    book = Workbook(); sheet = book.active; sheet.title = 'Revenue'
    sheet.append(['Spreadsheet preview', 'Amount']); sheet.append(['Revenue', 30]); sheet.append(['Cost', 12]); sheet.append(['Formula total', '=SUM(B2:B3)'])
    sheet.column_dimensions['A'].width = 30; sheet.column_dimensions['B'].width = 20
    other = book.create_sheet('Details'); other.append(['Spreadsheet second sheet', 42])
    other.column_dimensions['A'].width = 40
    book.save(DEST / 'matrix-sheet.xlsx')
    for name, mime, default in [('matrix-sheet.xlsx', 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet', True),
                              (convert('matrix-sheet.xlsx', 'xls:MS Excel 97'), 'application/vnd.ms-excel', True),
                              (convert('matrix-sheet.xlsx', 'ods'), 'application/vnd.oasis.opendocument.spreadsheet', False)]:
        add(name, mime, 'pdf', default, text='Spreadsheet preview', detail='42', secondText='Spreadsheet second sheet')

    deck = Presentation()
    for title in ['Presentation first slide', 'Presentation second slide']:
        slide = deck.slides.add_slide(deck.slide_layouts[5]); slide.shapes.title.text = title
        box = slide.shapes.add_textbox(SlideInches(1), SlideInches(2), SlideInches(6), SlideInches(1))
        box.text_frame.text = 'Talos slide body / 中文预览'
        slide.shapes.add_picture(str(DEST / 'matrix-image.png'), SlideInches(1), SlideInches(3), width=SlideInches(2))
    deck.save(DEST / 'matrix-slides.pptx')
    for name, mime, default in [('matrix-slides.pptx', 'application/vnd.openxmlformats-officedocument.presentationml.presentation', True),
                              (convert('matrix-slides.pptx', 'ppt:MS PowerPoint 97'), 'application/vnd.ms-powerpoint', False),
                              (convert('matrix-slides.pptx', 'odp'), 'application/vnd.oasis.opendocument.presentation', False)]:
        add(name, mime, 'pdf', default, text='Presentation first slide', secondText='Presentation second slide', chineseText='中文预览')

    for ext, mime, viewer, body in [
        ('md', 'text/markdown', 'markdown', '# Markdown format preview\n\n**Formatted body**\n\n| Item | Value |\n| --- | --- |\n| Total | 42 |'),
        ('csv', 'text/csv', 'text', 'item,value\nPreview total,42\n'),
        ('txt', 'text/plain', 'text', 'Plain text preview\n中文内容\n'),
        ('json', 'application/json', 'text', '{"preview":"JSON format","total":42}'),
        ('ts', 'text/plain', 'text', 'export const formatPreview = 42;\n'),
        ('xml', 'text/xml', 'text', '<?xml version="1.0"?><preview>XML format</preview>'),
        ('html', 'text/html', 'html', '<!doctype html><title>HTML format preview</title><h1>HTML format preview</h1><label>Draft<input aria-label="Format draft"></label>')]:
        name = 'matrix-content.' + ext; (DEST / name).write_text(body)
        add(name, mime, viewer, ext in ('md', 'csv'), text='Markdown format preview' if viewer == 'markdown' else 'HTML format preview' if viewer == 'html' else body.strip())

    member = b'Archive member preview: 42\n'
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, 'w', zipfile.ZIP_DEFLATED) as archive:
        archive.writestr('docs/readme.txt', member)
    (DEST / 'matrix-bundle.zip').write_bytes(buffer.getvalue())
    for suffix, mode in [('tar', 'w'), ('tgz', 'w:gz')]:
        with tarfile.open(DEST / ('matrix-bundle.' + suffix), mode) as archive:
            entry = tarfile.TarInfo('docs/readme.txt'); entry.size = len(member); entry.mtime = 0
            archive.addfile(entry, io.BytesIO(member))
    for ext, mime in [('zip', 'application/zip'), ('tar', 'application/x-tar'), ('tgz', 'application/gzip')]:
        add('matrix-bundle.' + ext, mime, 'archive', text=member.decode().strip())

    wav = DEST / 'matrix-tone.wav'
    with wave.open(str(wav), 'wb') as audio:
        audio.setparams((1, 2, 16000, 16000, 'NONE', 'not compressed'))
        audio.writeframes(b''.join(struct.pack('<h', int(1000 * math.sin(2 * math.pi * 440 * i / 16000))) for i in range(16000)))
    for ext, codec, mime in [('wav', None, 'audio/x-wav'), ('mp3', 'libmp3lame', 'audio/mpeg'), ('ogg', 'libvorbis', 'audio/ogg')]:
        name = 'matrix-tone.' + ext
        if codec:
            subprocess.run([ffmpeg, '-y', '-loglevel', 'error', '-i', str(wav), '-c:a', codec, str(DEST / name)], check=True)
        add(name, mime, 'audio', codec=codec or 'pcm_s16le')
    for ext, codec, mime in [('webm', 'libvpx', 'video/webm'), ('mp4', 'libx264', 'video/mp4')]:
        name = 'matrix-motion.' + ext
        subprocess.run([ffmpeg, '-y', '-loglevel', 'error', '-f', 'lavfi', '-i', 'testsrc2=size=160x96:rate=10', '-t', '1', '-c:v', codec, '-pix_fmt', 'yuv420p', '-an', str(DEST / name)], check=True)
        add(name, mime, 'video', codec='VP8' if ext == 'webm' else 'H.264', width=160, height=96)
    for case in cases:
        path = DEST / case['name']
        if path.stat().st_size > 500 * 1024:
            packed = path.with_name(path.name + '.gz')
            packed.write_bytes(gzip.compress(path.read_bytes(), mtime=0))
            case['fixture'] = packed.name
            path.unlink()
    (DEST / 'manifest.json').write_text(json.dumps(cases, indent=2, ensure_ascii=False) + '\n')
    print(f'{len(cases)} fixtures / {sum((DEST/c.get("fixture", c["name"])).stat().st_size for c in cases)} bytes')


if __name__ == '__main__':
    parser = argparse.ArgumentParser(); parser.add_argument('--ffmpeg', default='ffmpeg')
    generate(parser.parse_args().ffmpeg)
